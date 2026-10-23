"""Kernlogica: rapporten, planning, facturen en PDF-export."""
from __future__ import annotations

import base64
import csv
import hashlib
import hmac
import io
import re
import secrets
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import unquote

REPORT_EXTENSIONS = {".xls", ".xlsx", ".csv"}
ALLOWED_UPLOAD = REPORT_EXTENSIONS | {".pdf"}
PASSWORD_MIN_LENGTH = 4

CUSTOMER_ALIASES = {
    "powercast": "PowerCast",
    "huism": "Huisman",
    "huism.": "Huisman",
    "huisman": "Huisman",
    "verhoe": "Verhoeven",
    "verhoeven": "Verhoeven",
    "grooth": "Groothuis",
    "groothuis": "Groothuis",
    "werff": "Van der Werff",
    "vdwerff": "Van der Werff",
    "v.d.werff": "Van der Werff",
    "van der werff": "Van der Werff",
    "jverp": "J. van Erp",
    "j. van erp": "J. van Erp",
    "j van erp": "J. van Erp",
    "van erp": "J. van Erp",
    "loos": "Loos Kuijvegh",
    "loos kuijvegh": "Loos Kuijvegh",
    "loos-kuijvegh": "Loos Kuijvegh",
}


def validate_credentials(username: str, password: str) -> tuple[str, str]:
    username = (username or "").strip().lower()
    if not re.fullmatch(r"[a-z0-9._-]{3,40}", username):
        raise ValueError("Gebruikersnaam: 3-40 letters, cijfers, punten, streepjes of underscores.")
    if len(password or "") < PASSWORD_MIN_LENGTH:
        raise ValueError(f"Het wachtwoord moet minimaal {PASSWORD_MIN_LENGTH} tekens bevatten.")
    return username, password


def password_record(password: str, admin: bool = False) -> dict:
    salt = secrets.token_bytes(16)
    digest = hashlib.scrypt(password.encode("utf-8"), salt=salt, n=2**14, r=8, p=1, dklen=32)
    return {
        "salt": base64.b64encode(salt).decode("ascii"),
        "hash": base64.b64encode(digest).decode("ascii"),
        "admin": bool(admin),
        "created": datetime.now().isoformat(timespec="seconds"),
    }


def verify_password(password: str, record: dict) -> bool:
    try:
        salt = base64.b64decode(record["salt"])
        expected = base64.b64decode(record["hash"])
        actual = hashlib.scrypt(password.encode("utf-8"), salt=salt, n=2**14, r=8, p=1, dklen=32)
        return hmac.compare_digest(actual, expected)
    except (KeyError, ValueError, TypeError):
        return False


def public_users(users: dict) -> list[dict]:
    return [
        {"username": username, "admin": bool(record.get("admin")), "created": record.get("created") or ""}
        for username, record in sorted(users.items())
    ]


def money(value) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value or "").strip().replace(" ", "").replace("'", "")
    if not raw:
        return 0.0
    raw = raw.replace(",", ".")
    try:
        return float(raw)
    except ValueError:
        return 0.0


def normalize_plate(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", (value or "").upper())


def canon_customer(name: str) -> str:
    raw = " ".join((name or "").split())
    if not raw:
        return ""
    compact = re.sub(r"[.\s]+", "", raw).lower()
    return CUSTOMER_ALIASES.get(raw.lower()) or CUSTOMER_ALIASES.get(compact) or raw


def safe_filename(name: str) -> str:
    name = Path(unquote(name or "")).name
    name = re.sub(r"[^\w.\- ()]", "_", name, flags=re.UNICODE)
    return name.strip() or "upload.xls"


def is_week_name(name: str) -> bool:
    return name.lower().startswith("week")


def week_label(name: str) -> str:
    match = re.search(r"week\s*(\d{1,2})\s*[-_/ ]\s*(20\d{2})", name, re.I)
    if match:
        return f"Week {int(match.group(1))}-{match.group(2)}"
    return Path(name).stem or "Onbekende week"


def cell_text(value, datemode=None) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and datemode is not None:
        try:
            import xlrd

            if xlrd.xldate.xldate_as_tuple(value, datemode)[0] > 1900:
                dt = xlrd.xldate_as_datetime(value, datemode)
                if dt.hour or dt.minute or dt.second:
                    return dt.strftime("%H:%M:%S") if dt.year < 1901 else dt.strftime("%d/%m/%Y")
                return dt.strftime("%d/%m/%Y")
        except Exception:
            pass
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def iter_excel_sheets(path: Path, data: bytes | None = None):
    suffix = path.suffix.lower()
    if suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=data) if data is not None else xlrd.open_workbook(str(path))
        for sheet in workbook.sheets():
            rows = []
            for r in range(sheet.nrows):
                line = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        dt = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                        line.append(
                            dt.strftime("%H:%M:%S")
                            if dt.hour or dt.minute or dt.second
                            else dt.strftime("%d/%m/%Y")
                        )
                    else:
                        line.append(cell_text(cell.value))
                rows.append(line)
            yield sheet.name, rows
        return
    if suffix == ".xlsx":
        import openpyxl

        source = io.BytesIO(data) if data is not None else path
        workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
        for sheet in workbook.worksheets:
            rows = [[cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            yield sheet.title, rows
        workbook.close()


def list_report_files(uploaded: dict) -> list[dict]:
    files = []
    for name, item in uploaded.items():
        path = Path(name)
        if path.suffix.lower() not in REPORT_EXTENSIONS or item.get("kind") != "report":
            continue
        if path.stem.lower().startswith(("tolverdeling", "tol_per_klant")):
            continue
        files.append({"name": name, "size": len(item["data"]), "mtime": int(item["mtime"])})
    files.sort(key=lambda file: file["mtime"], reverse=True)
    return files


def pick(mapping: dict, *keys) -> str:
    lower = {str(k).strip().lower(): v for k, v in mapping.items() if k}
    for key in keys:
        val = mapping.get(key)
        if val not in (None, ""):
            return str(val).strip().strip("'")
        val = lower.get(key.lower())
        if val not in (None, ""):
            return str(val).strip().strip("'")
    return ""


def report_row_from_map(row: dict) -> dict | None:
    plate_raw = pick(row, "Truck PlateNr", "Card PlateNr", "File PlateNr", "Kenteken")
    plate = normalize_plate(plate_raw)
    if not plate:
        return None
    country = pick(row, "Trx Country", "Land")
    km = round(money(pick(row, "Distance KM", "KM")), 3)
    eur = round(money(pick(row, "Total Amt (EUR)", "Trx Amt (EUR)", "Tol EUR", "Tol")), 2)
    return {
        "id": pick(row, "Trx ID", "ID"),
        "date": pick(row, "Trx Date", "Datum"),
        "time": pick(row, "Trx Time", "Tijd"),
        "plate": plate,
        "plateRaw": plate_raw,
        "km": _correct_french_km(km, eur, country),
        "eur": eur,
        "road": pick(row, "Road Name", "Weg"),
        "entry": pick(row, "Entry Terminal"),
        "exit": pick(row, "Exit Terminal"),
        "tsp": pick(row, "TISPL TSP"),
        "country": country,
        "invId": pick(row, "Inv ID"),
        "invDate": pick(row, "Inv Date"),
        "spec": pick(row, "Spec Page"),
        "product": pick(row, "Prod Descr", "Product"),
        "billing": pick(row, "Billing Status"),
    }


def _correct_french_km(km: float, eur: float, country: str) -> float:
    """Telepass zet Franse TISPL-km een decimaal te weinig (12,45 i.p.v. 124,5)."""
    if km <= 0 or eur <= 0:
        return km
    code = re.sub(r"[^A-Z]", "", (country or "").upper())
    if code not in {"FR", "FRA", "FRANKRIJK", "FRANCE"}:
        return km
    if eur / km <= 1.2:
        return km
    fixed = round(km * 10, 3)
    if eur / fixed > 2.5:
        return km
    return fixed


def parse_csv_report(path: Path, data: bytes | None = None) -> list[dict]:
    rows = []
    handle = io.StringIO(data.decode("utf-8-sig"), newline="") if data is not None else path.open(encoding="utf-8-sig", newline="")
    with handle:
        reader = csv.DictReader(handle, delimiter=";")
        if not reader.fieldnames:
            return []
        for row in reader:
            parsed = report_row_from_map(row)
            if parsed:
                rows.append(parsed)
    return rows


def parse_excel_report(path: Path, data: bytes | None = None) -> list[dict]:
    rows = []
    for _name, sheet in iter_excel_sheets(path, data):
        header_idx = None
        headers = []
        for i, row in enumerate(sheet[:40]):
            labels = [str(c).strip() for c in row]
            joined = " ".join(labels).lower()
            if "truck platenr" in joined or "trx date" in joined or ("kenteken" in joined and "tol" in joined):
                header_idx = i
                headers = labels
                break
        if header_idx is None:
            continue
        for row in sheet[header_idx + 1 :]:
            mapping = {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))}
            parsed = report_row_from_map(mapping)
            if parsed:
                rows.append(parsed)
    return rows


def dmy_iso(dmy: str) -> str:
    raw = str(dmy or "").strip().strip("'\"")
    if " " in raw:
        raw = raw.split()[0]
    parts = [part for part in re.split(r"[/\-.]", raw) if part]
    if len(parts) != 3:
        return ""
    first, second, third = parts
    if len(first) == 4:
        year, month, day = first, second, third
    else:
        day, month, year = first, second, third
        if len(year) == 2:
            year = f"20{year}"
    try:
        return datetime(int(year), int(month), int(day)).strftime("%Y-%m-%d")
    except ValueError:
        return ""


def parse_report(name: str, invoice: str | None, uploaded: dict) -> list[dict]:
    item = uploaded.get(name)
    if not item:
        raise FileNotFoundError(name)
    path = Path(name)
    if path.suffix.lower() not in REPORT_EXTENSIONS:
        raise FileNotFoundError(name)
    data = item["data"]
    rows = parse_csv_report(path, data) if path.suffix.lower() == ".csv" else parse_excel_report(path, data)
    if invoice in (None, "", "alles"):
        return rows
    if invoice in ("ongefactureerd", "none"):
        return [r for r in rows if not r.get("invId")]
    return [r for r in rows if r.get("invId") == invoice]


def list_invoices(name: str, uploaded: dict) -> list[dict]:
    rows = dedupe_report_rows(parse_report(name, None, uploaded))
    buckets: dict[str, list] = {}
    for row in rows:
        buckets.setdefault(row.get("invId") or "ongefactureerd", []).append(row)
    out = []
    for inv_id, items in buckets.items():
        dates = [dmy_iso(r["date"]) for r in items if dmy_iso(r["date"])]
        out.append(
            {
                "id": inv_id,
                "date": items[0].get("invDate") or "",
                "from": min(dates) if dates else "",
                "to": max(dates) if dates else "",
                "rows": len(items),
                "eur": round(sum(r["eur"] for r in items), 2),
            }
        )
    out.sort(
        key=lambda x: (x["id"] != "ongefactureerd", dmy_iso(x["date"]) or x["from"] or ""),
        reverse=True,
    )
    return out


def extract_invoice_pdf_info(name: str, data: bytes) -> dict:
    info = {
        "name": name,
        "invoiceNumber": "",
        "invoiceDate": "",
        "customerNumber": "",
        "dutchToll": None,
        "total": None,
        "textReadable": False,
    }
    filename_match = re.search(r"(20\d{2})(\d{2})(\d{2})[_ -](\d+)", Path(name).stem)
    if filename_match:
        year, month, day, customer = filename_match.groups()
        info["invoiceDate"] = f"{year}-{month}-{day}"
        info["customerNumber"] = customer

    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:3])
    except Exception:
        text = ""
    compact = " ".join(text.split())
    if not compact:
        return info

    recognized = False
    for key, pattern in {
        "invoiceNumber": r"Factuurnummer\s*:?\s*(\d+)",
        "customerNumber": r"Klantnummer\s*:?\s*(\d+)",
    }.items():
        match = re.search(pattern, compact, re.IGNORECASE)
        if match:
            info[key] = match.group(1)
            recognized = True

    date_match = re.search(r"Factuurdatum\s*:?\s*(\d{1,2})[-/](\d{1,2})[-/](20\d{2})", compact, re.IGNORECASE)
    if date_match:
        day, month, year = date_match.groups()
        info["invoiceDate"] = f"{year}-{month.zfill(2)}-{day.zfill(2)}"
        recognized = True

    def amount(pattern: str) -> float | None:
        match = re.search(pattern, compact, re.IGNORECASE)
        if not match:
            return None
        return round(money(match.group(1).replace(".", "")), 2)

    info["dutchToll"] = amount(r"Nederlandse\s+tolwegen.*?(\d[\d.]*,\d{2})")
    info["total"] = amount(r"\bTotaal\s+(\d[\d.]*,\d{2})")
    info["textReadable"] = recognized or info["dutchToll"] is not None or info["total"] is not None
    return info


def list_invoice_pdf_info(uploaded: dict) -> list[dict]:
    return [
        extract_invoice_pdf_info(name, item["data"])
        for name, item in uploaded.items()
        if item.get("kind") == "invoice"
    ]


def check_invoice_pdf(report_name: str, invoice_id: str, uploaded: dict) -> dict:
    invoices = list_invoices(report_name, uploaded)
    invoice = next((item for item in invoices if item["id"] == invoice_id), None)
    if not invoice:
        raise ValueError("Factuur niet gevonden")
    pdfs = list_invoice_pdf_info(uploaded)
    if invoice_id == "ongefactureerd":
        return {
            "uploaded": len(pdfs),
            "match": None,
            "status": "none",
            "message": "Deze transacties zijn nog niet gefactureerd.",
        }

    target_date = dmy_iso(invoice.get("date") or "")
    match = next((pdf for pdf in pdfs if pdf["invoiceNumber"] == invoice_id), None)
    if not match and target_date:
        match = next((pdf for pdf in pdfs if pdf["invoiceDate"] == target_date), None)
    if not match:
        return {
            "uploaded": len(pdfs),
            "match": None,
            "status": "warning" if pdfs else "none",
            "message": "Geen bijpassende factuur-PDF geüpload." if pdfs else "",
        }

    amount_matches = match["dutchToll"] is None or abs(match["dutchToll"] - invoice["eur"]) < 0.01
    number_matches = not match["invoiceNumber"] or match["invoiceNumber"] == invoice_id
    date_matches = not target_date or match["invoiceDate"] == target_date
    verified = amount_matches and number_matches and date_matches
    if match["textReadable"]:
        message = (
            "Factuur-PDF sluit aan op het tolrapport."
            if verified
            else "Controleer deze PDF: niet alle factuurgegevens sluiten aan."
        )
    else:
        message = "Factuurdatum komt overeen. De gescande PDF bevat geen uitleesbare tekst voor verdere controle."
    return {
        "uploaded": len(pdfs),
        "match": match,
        "status": "ok" if verified else "warning",
        "message": message,
        "reportToll": invoice["eur"],
    }


def _sheet_date(title: str, header_cell: str) -> str | None:
    text = f"{title} {header_cell}"
    match = re.search(r"(\d{1,2})\s*-\s*(\d{1,2})\s*-?\s*(20\d{2})?", text)
    if not match:
        return None
    day, month, year = int(match.group(1)), int(match.group(2)), match.group(3)
    year = int(year) if year else datetime.now().year
    try:
        return datetime(year, month, day).date().isoformat()
    except ValueError:
        return None


def parse_week_sheet(name: str, sheet_rows: list[list[str]], trips: list, week: str) -> None:
    header = sheet_rows[0][3] if sheet_rows and len(sheet_rows[0]) > 3 else ""
    date_val = _sheet_date(name, header)
    last_customer = ""
    for row in sheet_rows:
        customer = (row[0] if row else "").strip()
        vehicle = (row[1] if len(row) > 1 else "").strip()
        plate_raw = (row[6] if len(row) > 6 else "").strip()
        trip_number = (row[2] if len(row) > 2 else "").strip()
        desc = (row[3] if len(row) > 3 else "").strip()
        km = money(row[5] if len(row) > 5 else 0)
        low = customer.lower()
        if low == "klant":
            continue
        if low.startswith("rit") or customer.replace(" ", "") in {"-", ".", "--"}:
            customer = last_customer
        elif customer:
            last_customer = customer
        else:
            customer = last_customer
        plate = normalize_plate(plate_raw)
        customer = canon_customer(customer)
        if not plate or not customer or not date_val:
            continue
        if customer.lower().startswith("loos") and not trip_number:
            trip_number = vehicle
        trips.append(
            {
                "week": week,
                "date": date_val,
                "plate": plate,
                "plateRaw": plate_raw,
                "customer": customer,
                "km": km,
                "tripNumber": trip_number or desc,
                "desc": desc,
            }
        )


def parse_week_files(uploaded: dict) -> dict:
    files = []
    trips = []
    sources = sorted(
        (
            (Path(name), item["data"])
            for name, item in uploaded.items()
            if item.get("kind") == "planning" and Path(name).suffix.lower() in {".xls", ".xlsx"}
        ),
        key=lambda item: item[0].name.lower(),
    )
    for path, data in sources:
        files.append(path.name)
        week = week_label(path.name)
        try:
            for sheet_name, rows in iter_excel_sheets(path, data):
                parse_week_sheet(sheet_name, rows, trips, week)
        except Exception as exc:
            return {"error": f"{path.name}: {exc}", "files": files, "rows": []}

    grouped: dict[tuple[str, str, str], dict] = {}
    for trip in trips:
        key = (trip["week"], trip["date"], trip["plate"])
        item = grouped.setdefault(
            key,
            {
                "week": trip["week"],
                "date": trip["date"],
                "plate": trip["plate"],
                "plateRaw": trip["plateRaw"],
                "km": 0.0,
                "customers": defaultdict(float),
                "tripNumbers": [],
                "desc": [],
            },
        )
        item["km"] += trip["km"]
        item["customers"][trip["customer"]] += trip["km"] or 1
        if trip["tripNumber"] and trip["tripNumber"] not in item["tripNumbers"]:
            item["tripNumbers"].append(trip["tripNumber"])
        if trip["desc"]:
            item["desc"].append(trip["desc"])
        if "-" in trip["plateRaw"]:
            item["plateRaw"] = trip["plateRaw"]

    rows = []
    for item in grouped.values():
        customer = max(item["customers"].items(), key=lambda kv: kv[1])[0]
        extra = [name for name in item["customers"] if name != customer]
        rows.append(
            {
                "week": item["week"],
                "date": item["date"],
                "plate": item["plate"],
                "plateRaw": item["plateRaw"],
                "customer": customer,
                "km": round(item["km"], 1),
                "tripNumber": " / ".join(item["tripNumbers"]),
                "desc": " / ".join(item["desc"][:3]),
                "extra": extra,
            }
        )
    rows.sort(key=lambda r: (r["week"], r["date"], r["plate"]))
    return {"files": files, "rows": rows}


def _trx_id_key(row: dict) -> str:
    return (row.get("id") or "").strip()


def _trx_fingerprint(row: dict) -> str:
    return "|".join(
        [
            row.get("date") or "",
            row.get("time") or "",
            row.get("plate") or "",
            f"{float(row.get('eur') or 0):.2f}",
            f"{float(row.get('km') or 0):.3f}",
        ]
    )


def _is_billed(row: dict) -> bool:
    return (row.get("billing") or "").strip().lower() == "billed"


def _is_dont_bill(row: dict) -> bool:
    return (row.get("billing") or "").strip().lower().replace("'", "") == "dont bill"


def _approx_num(left: float, right: float, tol: float) -> bool:
    return abs(left - right) <= tol


def _is_gantry_row(row: dict) -> bool:
    text = f"{row.get('road') or ''} {row.get('entry') or ''}"
    return bool(re.search(r"\b[AB]\s?\d{1,3}\b", text, re.I))


def _report_country_code(row: dict) -> str:
    raw = re.sub(r"[^A-Z]", "", (row.get("country") or "").upper())
    if raw in {"DE", "DEU", "GERMANY", "DEUTSCHLAND"} or raw.startswith("DEU"):
        return "DE"
    if len(raw) >= 2:
        return raw[:2]
    return raw


def _is_germany(row: dict) -> bool:
    return _report_country_code(row) == "DE"


_FOREIGN_TOLL = (
    "SANEF",
    "APRR",
    "COFIROUTE",
    "ALIS",
    "ASF",
    "AREA",
    "ESCOTA",
    "IBERPISTAS",
    "ALFAJARIN",
    "ASFINAG",
    "BRISA",
    "CONCESSAO",
    "AUTOSTRADE",
    "DARS",
    "VINCI",
    "MONTAUBAN",
    "CHESNES",
    "RIVES",
    "VALONGO",
    "INZERSDORF",
)


def _is_bundle_row(row: dict) -> bool:
    entry = (row.get("entry") or "").strip()
    km = float(row.get("km") or 0)
    if entry and re.fullmatch(r"[\d\s-]+", entry) and km >= 8:
        return True
    blob = f"{entry} {row.get('exit') or ''} {row.get('road') or ''} {row.get('tsp') or ''}".upper()
    if any(tag in blob for tag in _FOREIGN_TOLL):
        return True
    return (not _is_gantry_row(row)) and km >= 10


def collapse_detail_summaries(rows: list[dict]) -> list[dict]:
    """Duitse tol staat vaak als deeltrajecten én als somregel; tel de som niet dubbel."""
    groups: dict[tuple[str, str, str], list[int]] = defaultdict(list)
    for index, row in enumerate(rows):
        key = (row.get("date") or "", row.get("plate") or "", _report_country_code(row))
        groups[key].append(index)
    drop: set[int] = set()
    for indexes in groups.values():
        if not _is_germany(rows[indexes[0]]):
            continue
        gantry = [i for i in indexes if _is_gantry_row(rows[i])]
        if len(gantry) < 5:
            continue
        for i in indexes:
            if i in gantry:
                continue
            if _is_bundle_row(rows[i]):
                drop.add(i)
    for indexes in groups.values():
        if len(indexes) < 6:
            continue
        ordered = sorted(
            indexes,
            key=lambda i: (rows[i].get("time") or "", rows[i].get("id") or ""),
        )
        kms = {i: float(rows[i].get("km") or 0) for i in ordered}
        eurs = {i: float(rows[i].get("eur") or 0) for i in ordered}
        local_drop: set[int] = set()
        for i in sorted(ordered, key=lambda idx: kms[idx], reverse=True):
            if i in local_drop or i in drop:
                continue
            matched: list[int] = []
            acc_km = 0.0
            acc_eur = 0.0
            for j in reversed(ordered):
                if j == i or j in local_drop or j in drop:
                    continue
                if (rows[j].get("time") or "", rows[j].get("id") or "") >= (
                    rows[i].get("time") or "",
                    rows[i].get("id") or "",
                ):
                    continue
                if kms[j] >= kms[i] - 1e-9:
                    break
                acc_km += kms[j]
                acc_eur += eurs[j]
                matched.append(j)
                if acc_km > kms[i] + 0.08 and acc_eur > eurs[i] + 0.04:
                    break
                if (
                    len(matched) >= 5
                    and _approx_num(acc_km, kms[i], 0.08)
                    and _approx_num(acc_eur, eurs[i], 0.04)
                ):
                    local_drop.update(matched)
                    break
        drop.update(local_drop)
    return [row for i, row in enumerate(rows) if i not in drop]


def dedupe_report_rows(rows: list[dict]) -> list[dict]:
    """Zelfde rit niet twee keer: Don't Bill + Billed, of dezelfde Trx ID."""
    seen_ids: set[str] = set()
    by_fp: dict[str, dict] = {}
    for row in rows:
        trx_id = _trx_id_key(row)
        if trx_id:
            if trx_id in seen_ids:
                continue
            seen_ids.add(trx_id)
        fp = _trx_fingerprint(row)
        existing = by_fp.get(fp)
        if existing is None:
            by_fp[fp] = row
        elif _is_dont_bill(existing) and not _is_dont_bill(row):
            by_fp[fp] = row
    return collapse_detail_summaries(list(by_fp.values()))


def parse_all_reports(uploaded: dict) -> tuple[list[dict], list[str]]:
    """Alle geüploade tolrapporten in de sessie, gededupliceerd op transactie."""
    reports = list_report_files(uploaded)
    names = [item["name"] for item in reports]
    rows: list[dict] = []
    for name in names:
        rows.extend(parse_report(name, None, uploaded))
    return dedupe_report_rows(rows), names


def list_weeks(uploaded: dict) -> list[dict]:
    planning = parse_week_files(uploaded)
    if planning.get("error"):
        raise ValueError(planning["error"])
    buckets: dict[str, dict] = {}
    for row in planning.get("rows") or []:
        week = row["week"]
        item = buckets.setdefault(
            week,
            {"id": week, "from": row["date"], "to": row["date"], "trips": 0, "plates": set()},
        )
        item["from"] = min(item["from"], row["date"])
        item["to"] = max(item["to"], row["date"])
        item["trips"] += 1
        item["plates"].add(row["plate"])
    out = []
    for item in buckets.values():
        plates = item.pop("plates")
        item["vehicles"] = len(plates)
        out.append(item)
    out.sort(key=lambda week: (week["from"], week["id"]), reverse=True)
    return out


def detect_kind(path: Path, data: bytes | None = None) -> str:
    if path.suffix.lower() == ".pdf":
        return "invoice"
    if is_week_name(path.name):
        return "planning"
    if path.suffix.lower() == ".csv":
        try:
            if parse_csv_report(path, data):
                return "report"
        except Exception:
            pass
        return "unknown"
    try:
        for _name, rows in iter_excel_sheets(path, data):
            sample = " ".join(" ".join(r[:8]) for r in rows[:8]).lower()
            if "truck platenr" in sample or "trx date" in sample:
                return "report"
            if "klant" in sample and "kenteken" in sample:
                return "planning"
        return "unknown"
    except Exception:
        return "unknown"


def normalize_country(value: str) -> str:
    raw = " ".join((value or "").split())
    if not raw:
        return "Onbekend land"
    key = re.sub(r"[^a-z]", "", raw.lower())
    aliases = {
        "nl": "Nederland",
        "nld": "Nederland",
        "nederland": "Nederland",
        "netherlands": "Nederland",
        "holland": "Nederland",
        "be": "België",
        "bel": "België",
        "belgie": "België",
        "belgië": "België",
        "belgium": "België",
        "de": "Duitsland",
        "deu": "Duitsland",
        "duitsland": "Duitsland",
        "germany": "Duitsland",
        "deutschland": "Duitsland",
        "fr": "Frankrijk",
        "fra": "Frankrijk",
        "frankrijk": "Frankrijk",
        "france": "Frankrijk",
        "lu": "Luxemburg",
        "lux": "Luxemburg",
        "luxemburg": "Luxemburg",
        "luxembourg": "Luxemburg",
        "at": "Oostenrijk",
        "aut": "Oostenrijk",
        "oostenrijk": "Oostenrijk",
        "austria": "Oostenrijk",
        "ch": "Zwitserland",
        "che": "Zwitserland",
        "zwitserland": "Zwitserland",
        "switzerland": "Zwitserland",
        "pl": "Polen",
        "pol": "Polen",
        "polen": "Polen",
        "poland": "Polen",
        "cz": "Tsjechië",
        "cze": "Tsjechië",
        "tsjechie": "Tsjechië",
        "tsjechië": "Tsjechië",
        "czechrepublic": "Tsjechië",
        "it": "Italië",
        "ita": "Italië",
        "italie": "Italië",
        "italië": "Italië",
        "italy": "Italië",
        "es": "Spanje",
        "esp": "Spanje",
        "spanje": "Spanje",
        "spain": "Spanje",
        "dk": "Denemarken",
        "dnk": "Denemarken",
        "denemarken": "Denemarken",
        "denmark": "Denemarken",
    }
    return aliases.get(key, raw)


_TUNNEL_LABELS = (
    ("liefkenshoek", "Liefkenshoektunnel"),
    ("blankenburg", "Blankenburgtunnel"),
    ("westerschelde", "Westerscheldetunnel"),
    ("kennedytunnel", "Kennedytunnel"),
    ("coentunnel", "Coentunnel"),
    ("beneluxtunnel", "Beneluxtunnel"),
    ("sluiskil", "Sluiskiltunnel"),
    ("drechttunnel", "Drechttunnel"),
    ("heinenoord", "Heinenoordtunnel"),
    ("kiltunnel", "Kiltunnel"),
    ("vlaardingen", "Vlaardingentunnel"),
    ("wijkertunnel", "Wijkertunnel"),
    ("zeeburgertunnel", "Zeeburgertunnel"),
    ("noordzuid", "Noord/Zuid-Holland"),
)


def _tunnel_label(row: dict) -> str:
    """Tunnels staan in Spec Page / product / LHT, niet als gewone wegkilometers."""
    blob = " ".join(
        str(row.get(key) or "")
        for key in ("spec", "product", "road", "entry", "exit", "tsp")
    )
    compact = re.sub(r"[^a-z0-9]+", "", blob.lower())
    if "lht" in compact:
        return "Liefkenshoektunnel"
    for needle, label in _TUNNEL_LABELS:
        if needle in compact:
            return label
    if "tunnel" in compact:
        for key in ("spec", "product", "road"):
            raw = str(row.get(key) or "").strip()
            if raw and "tunnel" in raw.lower():
                return raw.split(";")[0].strip()
    return ""


def _split_amount(total: float, weights: list[float]) -> list[float]:
    """Verdeel total over weights; laatste regel vangt afronding."""
    if not weights:
        return []
    mass = sum(weights)
    if mass <= 0:
        share = round(total / len(weights), 1)
        parts = [share] * (len(weights) - 1)
        parts.append(round(total - share * (len(weights) - 1), 1))
        return parts
    parts = []
    assigned = 0.0
    for i, weight in enumerate(weights):
        if i == len(weights) - 1:
            parts.append(round(total - assigned, 1))
        else:
            share = round(total * weight / mass, 1)
            parts.append(share)
            assigned += share
    return parts


def _allocate_plan_km_by_country(rows: list[dict]) -> None:
    """Rittenboek heeft alleen totaal-km. Zet KM rit op het land, niet de hele rit.

    Betrouwbare landen (niet Frankrijk) krijgen hun KM tol als km in dat land.
    Frankrijk krijgt de rest van het rittenboek, zodat een Franse regel niet 760 km
    toont terwijl dat NL+BE+FR samen is.
    """
    trips: dict[tuple, list[dict]] = defaultdict(list)
    for row in rows:
        trips[(row.get("date"), row.get("plate"), row.get("customer"))].append(row)

    for group in trips.values():
        plan = float(group[0].get("planKm") or 0)
        if plan <= 0:
            for row in group:
                row["planKm"] = 0.0
            continue
        if len(group) == 1:
            continue
        unreliable = [row for row in group if (row.get("country") or "") == "Frankrijk"]
        reliable = [row for row in group if (row.get("country") or "") != "Frankrijk"]
        reliable_km = sum(float(row.get("km") or 0) for row in reliable)

        if unreliable and plan > reliable_km:
            leftover = round(plan - reliable_km, 1)
            for row in reliable:
                row["planKm"] = round(float(row.get("km") or 0), 1)
            shares = _split_amount(leftover, [float(row.get("km") or 0) for row in unreliable])
            for row, share in zip(unreliable, shares):
                row["planKm"] = share
        else:
            shares = _split_amount(plan, [float(row.get("km") or 0) for row in group])
            for row, share in zip(group, shares):
                row["planKm"] = share


def _french_toll_km_incomplete(plan_km: float, toll_km: float) -> bool:
    """Tolrapport mist Franse km als het duidelijk onder de rit uit het rittenboek blijft."""
    if plan_km < 80:
        return False
    return toll_km <= 0 or toll_km < 0.55 * plan_km


def _fill_incomplete_french_km(rows: list[dict]) -> None:
    """Frankrijk: onvolledige KM tol aanvullen uit het rittenboek. Bedrag blijft uit het rapport."""
    for row in rows:
        if (row.get("country") or "") != "Frankrijk":
            continue
        plan = float(row.get("planKm") or 0)
        km = float(row.get("km") or 0)
        if not _french_toll_km_incomplete(plan, km):
            continue
        row["kmReport"] = km
        row["km"] = plan
        row["kmFromPlan"] = True


def _resolve_plan(
    iso_date: str,
    plate: str,
    plan_map: dict,
    plan_by_plate: dict[str, list[dict]],
    week_from: str,
    week_to: str,
) -> tuple[dict | None, bool]:
    """Eerst exacte dag+kenteken, anders dezelfde auto elders in deze week (lege km / extra dag)."""
    exact = plan_map.get((iso_date, plate))
    if exact:
        return exact, True
    if not iso_date or not (week_from <= iso_date <= week_to):
        return None, False
    candidates = plan_by_plate.get(plate) or []
    if not candidates:
        return None, False
    target = datetime.strptime(iso_date, "%Y-%m-%d").date()
    closest = min(
        candidates,
        key=lambda row: abs((datetime.strptime(row["date"], "%Y-%m-%d").date() - target).days),
    )
    return closest, False


def _iso_week_label(iso_date: str) -> str:
    try:
        day = datetime.strptime(iso_date, "%Y-%m-%d").date()
    except ValueError:
        return ""
    iso = day.isocalendar()
    return f"Week {iso.week}-{iso.year}"


def _plan_index(plan_rows: list[dict]) -> dict:
    """Index over alle weekbestanden samen, zodat een tolregel maar één week krijgt."""
    exact: dict[tuple[str, str], dict] = {}
    by_week_plate: dict[tuple[str, str], list[dict]] = defaultdict(list)
    bounds: dict[str, dict] = {}
    for row in plan_rows:
        week = row["week"]
        item = bounds.setdefault(week, {"from": row["date"], "to": row["date"]})
        item["from"] = min(item["from"], row["date"])
        item["to"] = max(item["to"], row["date"])
        by_week_plate[(week, row["plate"])].append(row)
        key = (row["date"], row["plate"])
        current = exact.get(key)
        if current is None:
            exact[key] = row
            continue
        iso_label = _iso_week_label(row["date"])
        if row["week"] == iso_label and current["week"] != iso_label:
            exact[key] = row
        elif current["week"] != iso_label and row["week"] < current["week"]:
            exact[key] = row
    return {"exact": exact, "by_week_plate": by_week_plate, "bounds": bounds}


def _assign_plan(iso_date: str, plate: str, index: dict) -> tuple[dict | None, bool]:
    """Eén rit, één week: exacte dag wint; anders de dichtstbijzijnde rit in precies één week."""
    exact = (index.get("exact") or {}).get((iso_date, plate))
    if exact:
        return exact, True
    if not iso_date:
        return None, False
    try:
        target = datetime.strptime(iso_date, "%Y-%m-%d").date()
    except ValueError:
        return None, False
    best: dict | None = None
    best_key = None
    for week, bounds in (index.get("bounds") or {}).items():
        if not (bounds["from"] <= iso_date <= bounds["to"]):
            continue
        rows = (index.get("by_week_plate") or {}).get((week, plate)) or []
        if not rows:
            continue
        closest = min(
            rows,
            key=lambda row: abs((datetime.strptime(row["date"], "%Y-%m-%d").date() - target).days),
        )
        delta = abs((datetime.strptime(closest["date"], "%Y-%m-%d").date() - target).days)
        key = (delta, week)
        if best_key is None or key < best_key:
            best_key = key
            best = closest
    return best, False


def build_customer_pdf_rows(
    week_id: str, uploaded: dict
) -> tuple[dict, dict[str, list[dict]]]:
    """Rittenboek is leidend: overzicht per klant in die week. Tol komt uit alle rapporten."""
    _reject_unusable_uploads(uploaded)
    planning = parse_week_files(uploaded)
    if planning.get("error"):
        raise ValueError(planning["error"])
    all_plan_rows = planning.get("rows") or []
    plan_rows = [row for row in all_plan_rows if row["week"] == week_id]
    if not plan_rows:
        raise ValueError("Week niet gevonden in het rittenboek")

    index = _plan_index(all_plan_rows)
    week_from = min(row["date"] for row in plan_rows)
    week_to = max(row["date"] for row in plan_rows)
    transactions, report_names = parse_all_reports(uploaded)

    grouped: dict[tuple[str, str, str, str], dict] = {}
    matched_keys: set[tuple[str, str]] = set()
    unmatched_rows = 0
    unmatched_eur = 0.0

    for transaction in transactions:
        iso_date = dmy_iso(transaction.get("date") or "")
        plate = transaction["plate"]
        plan, exact = _assign_plan(iso_date, plate, index)
        if not plan or plan.get("week") != week_id:
            if not plan and iso_date and week_from <= iso_date <= week_to:
                unmatched_rows += 1
                unmatched_eur += transaction["eur"]
            continue
        matched_keys.add((iso_date, plate))
        country = normalize_country(transaction.get("country") or "")
        tunnel = _tunnel_label(transaction)
        customer = plan["customer"]
        key = (iso_date, plate, customer, country)
        item = grouped.setdefault(
            key,
            {
                "date": datetime.strptime(iso_date, "%Y-%m-%d").date() if iso_date else None,
                "plate": plan.get("plateRaw") or transaction.get("plateRaw") or plate,
                "ritten": 0,
                "km": 0.0,
                "planKm": round(float(plan.get("km") or 0), 1) if exact else 0.0,
                "eur": 0.0,
                "billedEur": 0.0,
                "openEur": 0.0,
                "tunnelEur": 0.0,
                "tunnels": [],
                "customer": customer,
                "country": country,
                "ritnr": plan.get("tripNumber") or "",
            },
        )
        item["ritten"] += 1
        if tunnel:
            item["tunnelEur"] += transaction["eur"]
            if tunnel not in item["tunnels"]:
                item["tunnels"].append(tunnel)
        else:
            item["km"] += transaction["km"]
            item["eur"] += transaction["eur"]
        if _is_billed(transaction):
            item["billedEur"] += transaction["eur"]
        elif _is_dont_bill(transaction):
            item["openEur"] += transaction["eur"]
        if "-" in (transaction.get("plateRaw") or ""):
            item["plate"] = transaction["plateRaw"]

    billed_items = []
    for item in grouped.values():
        item["km"] = round(item["km"], 1)
        item["eur"] = round(item["eur"], 2)
        item["billedEur"] = round(item["billedEur"], 2)
        item["openEur"] = round(item["openEur"], 2)
        item["tunnelEur"] = round(item.get("tunnelEur") or 0, 2)
        item["tunnel"] = " + ".join(item.get("tunnels") or [])
        if item["eur"] <= 0 and item["tunnelEur"] <= 0:
            continue
        billed_items.append(item)
    _allocate_plan_km_by_country(billed_items)
    _fill_incomplete_french_km(billed_items)

    by_customer: dict[str, list[dict]] = defaultdict(list)
    for item in billed_items:
        by_customer[item["customer"]].append(item)
    for rows in by_customer.values():
        rows.sort(key=lambda row: ((row["date"] or date.min), row.get("country") or "", row["plate"]))

    meta = {
        "id": week_id,
        "from": week_from,
        "to": week_to,
        "reports": report_names,
        "planningTrips": len(plan_rows),
        "unmatchedTollRows": unmatched_rows,
        "unmatchedTollEur": round(unmatched_eur, 2),
        "qualityNotes": _quality_notes(unmatched_rows, unmatched_eur, dict(by_customer)),
    }
    return meta, dict(by_customer)


def _quality_notes(unmatched_rows: int, unmatched_eur: float, by_customer: dict) -> list[dict]:
    """Controles voor de operator bij elke nieuwe week; niet op de klant-PDF."""
    notes = []
    if unmatched_eur > 1:
        notes.append(
            {
                "level": "warning",
                "text": (
                    f"{unmatched_rows} tolregels in deze week zitten niet in het rittenboek "
                    f"(€ {unmatched_eur:,.2f}). Die komen niet op een klant-PDF."
                ),
            }
        )
    filled_fr = sum(
        1
        for rows in by_customer.values()
        for row in rows
        if row.get("kmFromPlan")
    )
    if filled_fr:
        notes.append(
            {
                "level": "info",
                "text": (
                    f"{filled_fr} Franse rit(ten): kilometers aangevuld uit het rittenboek "
                    "(het tolrapport had die km niet volledig). Het bedrag blijft uit het rapport."
                ),
            }
        )
    return notes


def _week_share_of_report(
    week_id: str,
    plan_rows: list[dict],
    date_counts: Counter,
    total_rows: int,
) -> dict:
    """Welk deel van alle tolregels bij deze weekplanning hoort."""
    week_dates = {row["date"] for row in plan_rows if row.get("date")}
    covered = sum(date_counts[day] for day in week_dates)
    percent = round(100 * covered / total_rows) if total_rows else 0
    return {
        "id": week_id,
        "trips": len(plan_rows),
        "days": len(week_dates),
        "covered": covered,
        "percent": percent,
        "from": min(week_dates) if week_dates else "",
        "to": max(week_dates) if week_dates else "",
        "message": (
            f"{covered} van {total_rows} tolregels ({percent}%) horen bij {week_id}"
            + (f" ({min(week_dates)} t/m {max(week_dates)})." if week_dates else ".")
        ),
    }


def _collapse_day_ranges(days: list[str]) -> list[dict]:
    parsed = sorted(datetime.strptime(day, "%Y-%m-%d").date() for day in days)
    if not parsed:
        return []
    ranges = []
    start = prev = parsed[0]
    for day in parsed[1:]:
        if day == prev + timedelta(days=1):
            prev = day
            continue
        ranges.append({"from": start.isoformat(), "to": prev.isoformat()})
        start = prev = day
    ranges.append({"from": start.isoformat(), "to": prev.isoformat()})
    return ranges


def _calendar_days(
    report_from: str,
    report_to: str,
    plan_rows: list[dict],
    date_counts: Counter,
) -> list[dict]:
    week_by_date: dict[str, str] = {}
    for row in plan_rows:
        day = row.get("date") or ""
        if day and day not in week_by_date:
            week_by_date[day] = row.get("week") or ""
    start = datetime.strptime(report_from, "%Y-%m-%d").date()
    end = datetime.strptime(report_to, "%Y-%m-%d").date()
    days = []
    cursor = start
    while cursor <= end:
        iso = cursor.isoformat()
        days.append(
            {
                "date": iso,
                "inPlanning": iso in week_by_date,
                "inReport": iso in date_counts,
                "rows": int(date_counts.get(iso, 0)),
                "week": week_by_date.get(iso, ""),
            }
        )
        cursor += timedelta(days=1)
    return days


def _missing_report_ranges(report_from: str, report_to: str, plan_dates: set[str], date_counts: Counter) -> list[dict]:
    start = datetime.strptime(report_from, "%Y-%m-%d").date()
    end = datetime.strptime(report_to, "%Y-%m-%d").date()
    missing_days = []
    cursor = start
    while cursor <= end:
        iso = cursor.isoformat()
        if iso not in plan_dates:
            missing_days.append(iso)
        cursor += timedelta(days=1)
    ranges = _collapse_day_ranges(missing_days)
    for item in ranges:
        item["rows"] = sum(count for day, count in date_counts.items() if item["from"] <= day <= item["to"])
    return ranges


def check_reports_against_planning(uploaded: dict) -> dict:
    """Dekking van alle tolrapporten tegen alle geüploade weekplanningen samen."""
    planning = parse_week_files(uploaded)
    if planning.get("error"):
        raise ValueError(planning["error"])
    plan_rows = planning.get("rows") or []
    week_ids = sorted({row["week"] for row in plan_rows})
    plan_dates = {row["date"] for row in plan_rows}
    plan_keys = {(row["date"], row["plate"]) for row in plan_rows}
    report_files = list_report_files(uploaded)

    empty = {
        "weeks": week_ids,
        "planningFrom": min(plan_dates) if plan_dates else "",
        "planningTo": max(plan_dates) if plan_dates else "",
        "planningDays": len(plan_dates),
        "planningTrips": len(plan_rows),
        "reports": [],
        "totals": {
            "status": "warning",
            "message": "Upload tolrapporten én weekplanningen om te zien hoeveel je nog mist.",
            "rows": 0,
            "covered": 0,
            "missing": 0,
            "matched": 0,
            "percent": 0,
        },
        "planning": {"trips": 0, "days": 0, "covered": 0, "percent": 0, "message": ""},
        "weeksDetail": [],
        "timeline": {"reportFrom": "", "reportTo": "", "missing": [], "days": []},
    }
    if not plan_rows or not report_files:
        return empty

    checks = []
    total_rows = 0
    total_covered = 0
    total_matched = 0
    report_keys: set[tuple[str, str]] = set()
    date_counts: Counter = Counter()
    week_names = ", ".join(week_ids)
    plan_span = f"{min(plan_dates)} t/m {max(plan_dates)}"

    for file_info in report_files:
        name = file_info["name"]
        dated = []
        covered = 0
        matched = 0
        for row in dedupe_report_rows(parse_report(name, None, uploaded)):
            iso = dmy_iso(row.get("date") or "")
            if not iso:
                continue
            dated.append(iso)
            date_counts[iso] += 1
            report_keys.add((iso, row["plate"]))
            if iso in plan_dates:
                covered += 1
                if (iso, row["plate"]) in plan_keys:
                    matched += 1
        total = len(dated)
        missing = total - covered
        pct = round(100 * covered / total) if total else 0
        if total == 0:
            status = "warning"
            message = f"{name}: geen transactiedatums gevonden."
        elif covered == 0:
            status = "mismatch"
            message = (
                f"{name}: 0 van {total} regels vallen in de geüploade weekplanningen ({week_names}, {plan_span}). "
                f"Nog {missing} te gaan. Tolperiode: {min(dated)} t/m {max(dated)}."
            )
        elif missing == 0:
            status = "ok"
            message = (
                f"{name}: alle {total} regels vallen in de geüploade weekplanningen ({week_names}). "
                f"{matched} gekoppeld op kenteken."
            )
        else:
            status = "ok" if pct >= 85 else "warning"
            message = (
                f"{name}: {covered} van {total} regels ({pct}%) vallen in de geüploade weekplanningen "
                f"({week_names}, {plan_span}). Nog {missing} regels missen een weekplanning. "
                f"Tolperiode: {min(dated)} t/m {max(dated)}."
            )
        checks.append(
            {
                "name": name,
                "status": status,
                "message": message,
                "from": min(dated) if dated else "",
                "to": max(dated) if dated else "",
                "rows": total,
                "covered": covered,
                "missing": missing,
                "matched": matched,
                "percent": pct,
            }
        )
        total_rows += total
        total_covered += covered
        total_matched += matched

    missing = total_rows - total_covered
    pct = round(100 * total_covered / total_rows) if total_rows else 0
    if missing == 0:
        totals_status = "ok"
        totals_message = (
            f"Alle {total_rows} tolregels vallen in de geüploade weekplanningen ({week_names}). "
            f"{total_matched} gekoppeld op kenteken."
        )
    elif total_covered == 0:
        totals_status = "mismatch"
        totals_message = (
            f"Nog niets gedekt: 0 van {total_rows} tolregels vallen in {week_names} ({plan_span}). "
            f"Upload de weekplanningen die bij de tolperiode horen."
        )
    else:
        totals_status = "ok" if pct >= 85 else "warning"
        totals_message = (
            f"Samen: {total_covered} van {total_rows} tolregels ({pct}%) vallen in de geüploade weekplanningen "
            f"({week_names}, {plan_span}). Nog {missing} regels missen een weekplanning."
        )

    return {
        "weeks": week_ids,
        "planningFrom": min(plan_dates),
        "planningTo": max(plan_dates),
        "planningDays": len(plan_dates),
        "planningTrips": len(plan_rows),
        "reports": checks,
        "totals": {
            "status": totals_status,
            "message": totals_message,
            "rows": total_rows,
            "covered": total_covered,
            "missing": missing,
            "matched": total_matched,
            "percent": pct,
        },
        "weeksDetail": [
            _week_share_of_report(
                week_id,
                [row for row in plan_rows if row["week"] == week_id],
                date_counts,
                total_rows,
            )
            for week_id in dict.fromkeys(week_ids)
        ],
        "timeline": {
            "reportFrom": min(date_counts) if date_counts else "",
            "reportTo": max(date_counts) if date_counts else "",
            "missing": _missing_report_ranges(min(date_counts), max(date_counts), plan_dates, date_counts)
            if date_counts
            else [],
            "days": _calendar_days(min(date_counts), max(date_counts), plan_rows, date_counts)
            if date_counts
            else [],
        },
    }


def check_reports_for_week(week_id: str, uploaded: dict, plan_rows: list[dict] | None = None) -> list[dict]:
    """Blijft beschikbaar: dekt alle planningen, genegeerde week_id."""
    return check_reports_against_planning(uploaded).get("reports") or []


def _rows_for_customer(by_customer: dict[str, list[dict]], name: str) -> list[dict] | None:
    wanted = canon_customer(name)
    if wanted in by_customer:
        return by_customer[wanted]
    low = wanted.lower()
    compact = re.sub(r"[.\s]+", "", wanted).lower()
    for key, rows in by_customer.items():
        if key.lower() == low or re.sub(r"[.\s]+", "", key).lower() == compact:
            return rows
        if canon_customer(key).lower() == low:
            return rows
    return None


def _reject_unusable_uploads(uploaded: dict) -> None:
    huge = [name for name, item in uploaded.items() if item.get("too_large")]
    if huge:
        raise ValueError(
            "Het tolrapport in deze sessie is te groot voor de website. "
            "Klik Wissen, upload eerst het rittenboek en daarna opnieuw het CSV."
        )


def list_pdf_customers(week_id: str, uploaded: dict) -> dict:
    _reject_unusable_uploads(uploaded)
    meta, by_customer = build_customer_pdf_rows(week_id, uploaded)
    merged: dict[str, dict] = {}
    for customer, rows in by_customer.items():
        eur = round(sum(row["eur"] + float(row.get("tunnelEur") or 0) for row in rows), 2)
        if eur <= 0:
            continue
        customer = canon_customer(customer)
        key = customer.lower()
        item = merged.setdefault(
            key,
            {
                "customer": customer,
                "countries": set(),
                "rows": 0,
                "trips": 0,
                "km": 0.0,
                "eur": 0.0,
            },
        )
        if len(customer) > len(item["customer"]):
            item["customer"] = customer
        item["countries"].update(row["country"] for row in rows if row.get("country"))
        item["rows"] += len(rows)
        item["trips"] += sum(row["ritten"] for row in rows)
        item["km"] += sum(row["km"] for row in rows)
        item["eur"] += eur
    customers = []
    for item in sorted(merged.values(), key=lambda row: row["customer"].lower()):
        countries = sorted(item["countries"])
        customers.append(
            {
                "customer": item["customer"],
                "country": ", ".join(countries),
                "countries": countries,
                "rows": item["rows"],
                "trips": item["trips"],
                "km": round(item["km"], 1),
                "eur": round(item["eur"], 2),
            }
        )
    return {
        **meta,
        "customers": customers,
    }


def _json_ready(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(key): _json_ready(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_json_ready(item) for item in value]
    if isinstance(value, float) and (value != value or abs(value) == float("inf")):
        return 0.0
    return value


def build_preview_cache(uploaded: dict) -> dict:
    """Lichte JSON voor Vercel GET /api/weeks en /api/pdf-customers."""
    weeks = list_weeks(uploaded)
    pdf = {}
    for week in weeks:
        pdf[week["id"]] = list_pdf_customers(week["id"], uploaded)
    planning = [
        name
        for name, item in uploaded.items()
        if item.get("kind") == "planning"
    ]
    return _json_ready(
        {
            "weeks": weeks,
            "pdf": pdf,
            "planningFiles": sorted(planning),
            "reports": list_report_files(uploaded),
        }
    )


def build_customer_pdf(week_id: str, customer: str, country: str, uploaded: dict) -> bytes:
    import maak_pdf_per_klant as pdf

    meta, by_customer = build_customer_pdf_rows(week_id, uploaded)
    rows = _rows_for_customer(by_customer, customer)
    if not rows:
        raise ValueError("Klant niet gevonden in deze week van het rittenboek")
    if country:
        wanted = normalize_country(country)
        rows = [row for row in rows if row.get("country") == wanted]
        if not rows:
            raise ValueError("Deze klant heeft in deze week geen tol in dat land")
    source = ", ".join(meta.get("reports") or []) or "rittenboek"
    return pdf.make_pdf_bytes(customer, rows, meta, source)


def build_customer_zip(week_id: str | None, uploaded: dict) -> bytes:
    import maak_pdf_per_klant as pdf

    weeks = list_weeks(uploaded)
    if week_id and week_id != "alles":
        weeks = [item for item in weeks if item["id"] == week_id]
    if not weeks:
        raise ValueError("Geen week gevonden in het rittenboek")

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for week in weeks:
            current, by_customer = build_customer_pdf_rows(week["id"], uploaded)
            folder = pdf.safe_name(current["id"])
            source = ", ".join(current.get("reports") or []) or "rittenboek"
            for customer, rows in sorted(by_customer.items(), key=lambda item: item[0].lower()):
                if round(sum(row["eur"] + float(row.get("tunnelEur") or 0) for row in rows), 2) <= 0:
                    continue
                filename = f"Toloverzicht - {pdf.safe_name(current['id'])} - {pdf.safe_name(customer)}.pdf"
                archive.writestr(
                    f"{folder}/{filename}",
                    pdf.make_pdf_bytes(customer, rows, current, source),
                )
    return buffer.getvalue()
